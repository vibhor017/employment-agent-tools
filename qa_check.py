#!/usr/bin/env python3
"""
qa_check.py, mechanical QA gate for the Employment Agent pipeline.

Usage:
  python3 qa_check.py cl  <file> [--jd JD.txt]   # cover letter (.docx, .md or .txt)
  python3 qa_check.py cv  <file> [--jd JD.txt]   # CV
  python3 qa_check.py net <file>                 # networking package or interview prep
  python3 qa_check.py hm  <file>                 # Hiring Manager Review record (hm_review.md), run before cl/cv

Exit code 0 = PASS, 1 = FAIL (blocking issues found). Warnings never block.
Checks only what a regex can check. Judgment items (the argument spine, voice,
flow, bridge honesty, whether a hook or vignette is true to its Story Bank entry) belong to
the Hiring Manager Review, APPLICATION_RULES.md Section 8A, which runs before
this script.

MAINTENANCE RULE: this file mirrors CANDIDATE_STRATEGY.md Section 5B and
APPLICATION_RULES.md Sections 0, 0.2a, 2, 2.5, 11 and 12.6. When either changes, update
BANNED_CLAIMS, RETIRED_FIGURES, BANNED_PHRASES, HARD_LEXICON, SOFT_LEXICON or
ENTRY_LINKS here in the same edit, or the gate falls behind the rules.

v1.7, 24 October 2026 (project instructions v9.6):
- New `hm` mode: checks the Hiring Manager Review record (APPLICATION_RULES.md 8A) that the agent writes to
  <app_dir>/hm_review.md before the first cl or cv run. FAIL unless items 1-12 each appear once, each with PASS or
  FAIL and a quoted line from the draft ("..." of 3+ words). WARN on a FAIL item not marked revised. Prints N/12 pass.
  build_application.js v1.1 refuses to build the .docx files until this file exists.

v1.6, 24 October 2026 (APPLICATION_RULES.md v12.4, CANDIDATE_STRATEGY.md v9.4):
- Product name: the product is Gyro Health. The old name "Curo" FAILS in cv and cl and WARNS in net; so does
  any "formerly called", "previously known as", "rebranded" or "renamed" next to the product. Every product
  pattern below matches both names. Switches renamed GYRO_COFOUNDER_ALLOWED and GYRO_PAGE_LIVE (both False);
  the page must carry the new name and description before GYRO_PAGE_LIVE turns True. The deep-link slug stays
  vibhor.work/project/curo in ENTRY_LINKS until the site moves the page.
- AI-isms lexicon (APPLICATION_RULES.md 0.2a). HARD_LEXICON: FAIL in cl and net, WARN in cv. SOFT_LEXICON: WARN
  everywhere. "harness" counts only as a verb ("harness the power", "harnessing"); the noun in "evaluation
  harness" is a real product term. "aligned to" is required wording and is never flagged.
- Curly apostrophes and quotes are normalised before every check.
- Cover letter voice (0.3a): WARN at zero contractions (B and C may leave it standing as "formal register") and
  above three; WARN when sentence lengths have a standard deviation under 5 words; WARN on "By [doing X], I".
- No question anywhere in a letter body (FAIL). Para 1 hook (2.1): FAIL on "Imagine", "Picture this",
  "In today's" or a rhetorical-question opener; WARN on an "As a" opener and on a first sentence over 25 words.
  FAIL on visa, residence-permit or work-permit wording in Para 1; WARN on it elsewhere (outreach only).
- Evidence-table letters (APPLICATION_RULES.md 2.5): rows read from the .docx table or from "| left | right |"
  lines in a .txt draft, written in the same block as the Para 3 lead-in and follow-up (no blank lines between).
  FAIL outside Para 3, fewer than 3 or more than 4 rows, a row without two cells, a right-hand cell with no
  numeral, a figure repeated across rows, a deep link inside a cell. WARN over 6 words left or 25 words right.
  Totals 300-360 words, cells counted. Budgets: Para 1 85-110, Para 2 95-125, Para 3 80-115, Para 4 35-55.
- Prose budgets: Para 1 95-120, Para 2 115-140, Para 3 105-125, Para 4 40-60 (total unchanged, 360-400).

v1.5, October 2026 (CURO_UPDATE_2026-10-21.md):
- Switches GYRO_COFOUNDER_ALLOWED and CURO_PAGE_LIVE (both False). While False, no co-founder, CTO or named
  teammate may appear in a CV or letter, and the vibhor.work/project/curo deep link fails in every mode.
- LEGACY_CURO (FAIL in cv and cl, WARN in net): OBGYN, Flare Mode, Curo Heroes, Van Westendorp, the old 15-patient
  cohort, a cohort that "set decisions before code", Article 50 overclaims (compliant, certified, sprint), 59 of 61 and 61 tasks, EU-region, Weekly Reciprocity, Day 21, 6-tap, five-event, PCOS or Germany near
  Curo, a named teammate, the Stipendium claim, a former co-founder.
- HELD_CURO (WARN): MDCG, DiGA, zero-gap audit claims, build counts. Article 50 is a current claim and is not held.
- COHORT (WARN): the word "cohort" with no pilot, planned, about 35 or enrol nearby (the planned pilot cohort is allowed).
- WARN: gate figures (at most 5%, at least 60%, 100 red-team prompts, zero escapes) without target, gate or plan
  nearby; a steelman or expert committee without "simulated".
- FAIL: a pilot described as run, completed or conducted (WARN in net). WARN: "pilot results", "interview findings".
- Founder disclosure now looks for the founder wording, not "co-founder".
- KEY_FIGURES: "59 of 61" and "15-patient" removed, "13 experiments" added. JARGON gains "case-crossover".

v1.4, October 2026 (Curo repo audit, claude/CURO_BULLET_BANK.md; IPAC detail):
- Retired figure: 60-70% false positive reduction (superseded by the repo's
  68% to 3% in simulation).
- New 5B FAILs: a CTO-led or co-built Curo build (Curo was built solo with AI
  coding agents), cohort "weekly feedback", a per-token cost model, the
  38-47 week estimate and the 16-week MVP scope (unsourced until Vibhor
  confirms the estimate), Charite BIH as a partner (unconfirmed), "shipped the
  MVP", the engine as live or in production, an ML-based output filter, and
  MoSCoW on Curo. Remove a pattern only when Vibhor confirms the claim, in the
  same edit as CANDIDATE_STRATEGY.md 5B.
- The "in simulation" check now covers the repo figures (68%, 19.5%, 33-67%):
  each needs "simulat" within its window.
- CONFIRMED_TOOLS gains Cursor, Claude Code, Expo, AWS Bedrock, Vertex AI and
  pgvector (stated by Vibhor or named in the Curo repo bank).
- KEY_FIGURES tracks 68%, 19.5%, 33-67%, 12 days and 59 of 61.
- Thesis title: "Operationalises Constitutional Equality" is retired wording; the
  title reads "...Operationalising Constitutional Equality".
v1.2, October 2026:
- Cover letter .docx parsing reads the body from the "Dear" line to the
  sign-off, so the letterhead no longer counts as a paragraph or as words.
  Paragraphs are joined with blank lines so the four-paragraph check works.
- Profile Summary checked on its own (55-word cap, no "I", three sentences);
  it no longer trips the 35-word bullet warning.
- New FAILs: retired figures (30% resolution, 200+ residents, 30+ advisory
  reports, 15-20 patient cohort, old thesis titles, years phrasing other than
  "5+ years"), CMGRS individual-ownership claims, the north star as the cause
  of the 28%, the GDPR "product changes" client claim, "Launched Curo", and a
  delivered 16-week MVP.
- New WARNs: Para 3 without an "I would" approach sentence; Para 4 not
  returning to Para 1's words; three consecutive sentences with no I or we in
  Para 1 or Para 3; contact details in Para 4; more than 5 Curo or 4 IPAC
  bullets; a figure used more than twice in the CV; method names in CV lines
  the JD does not use (optional --jd); the old date format and the old
  "Products Built" heading; Problem:/Solution: lines over 40 words.
v1.3, October 2026 (APPLICATION_RULES.md 11.3a and 11.3b):
- Areas of Expertise echo check (FAIL): every term, split into its parts at
  "and", "&", "/" and commas, needs one of that part's distinctive words (or its
  stem) in a bullet, Key Achievement, summary, certification or education line.
  Generic words (product, delivery, management, development, strategy,
  planning, design, process, framework, platform) never count when the part has a distinctive word,
  so "Agile Delivery" fails unless "agile" appears. Tool names inside brackets
  are skipped; tools belong on the Tools line.
- Tools line (FAIL): a "Tools:" line must be present, every tool on it must be
  in CONFIRMED_TOOLS, and skill-level words (proficient, expert, advanced...)
  are not allowed. WARN over 10 tools, and (with --jd) WARN when the JD names a
  confirmed tool that the line leaves out.
- Both checks read .docx (table cells and the "Tools:" paragraph) and .md/.txt
  (the lines between the AREAS OF EXPERTISE heading and the next section).
v1.2.1, October 2026: .docx CVs and letters must carry the LinkedIn profile and
the portfolio as real hyperlinks on the contact line (FAIL if either is plain
text or missing).
v1.1, September 2026: vibhor.work as the only domain; deep links validated
against the dossier entry list; 5B coverage extended.
"""
import re, sys, statistics

CANONICAL_DOMAIN = "vibhor.work"
# Switches (v1.5, renamed v1.6)
GYRO_COFOUNDER_ALLOWED = False   # True only if a co-founder or CTO is confirmed and may be named
GYRO_PAGE_LIVE = False           # True once the product page on vibhor.work carries the Gyro Health name and matches
                                 # CURO_UPDATE_2026-10-21.md Section 4 (update the slug in ENTRY_LINKS if the page moves)
PRODUCT = r"(?:Gyro Health|Curo)"   # every product pattern matches the current and the old name
OLD_NAME = re.compile(r"\bCuro\b")
RENAME_MENTION = re.compile(r"\b(?:formerly|previously)\s+(?:called|known as|named)\b|\brebrand(?:ed|ing)?\b|"
                            r"Gyro Health[^.]{0,60}\brenamed\b|\brenamed\b[^.]{0,60}Gyro Health", re.I)

FOUNDER_DISCLOSURE = re.compile(
    r"\bI(?: am|'m) (?:the )?founder\b|\bI founded\b|\bfound(?:er|ed)\b[^.]{0,80}\b(?:Gyro Health|Curo)\b|\b(?:Gyro Health|Curo)\b[^.]{0,80}\bfound(?:er|ed)\b",
    re.I)
RETIRED_DOMAIN = "producthq.workers.dev"   # same build, never used in artefacts
DEEP_LINK = re.compile(r"vibhor\.work/(?:project|case-study)/[a-z0-9-]+", re.I)
# PORTFOLIO_DOSSIER.md track routing table. Umang is never linked.
ENTRY_LINKS = {
    "vibhor.work/project/curo", "vibhor.work/project/cmgrs",
    "vibhor.work/project/rasahub", "vibhor.work/project/haqdaar",
    "vibhor.work/project/beamerheld", "vibhor.work/case-study/typecast",
    "vibhor.work/case-study/commonstrace",
}

# Section 0.1 dash ban
DASHES = {"—": "em dash", "–": "en dash"}

# Section 0.2 / 8.0 AI tells and denial sentences (blocking)
BANNED_PHRASES = [
    r"not a pivot", r"not a career change", r"despite (?:a|my) non-traditional",
    r"unlike most candidates", r"it is not just", r"it's not just",
    r"this is precisely why", r"it is worth noting", r"\bfundamentally\b",
    r"\bi am applying\b", r"\bi am excited\b", r"\bi'm excited\b",
    r"\bsuccessfully\b", r"\beffectively\b", r"\bproactively\b",
    r"side project\b(?!s)",
]

# APPLICATION_RULES.md Section 0.2a, the AI-isms lexicon.
# Tier 1: FAIL in letters and outreach (cl, net), WARN on CVs.
HARD_LEXICON = [
    # words
    r"\bdelv(?:e|es|ed|ing)\b", r"\brealms?\b", r"\bintricate\b", r"\bshowcas(?:e|es|ed|ing)\b",
    r"\bpivotal\b", r"\bleverag(?:e|es|ed|ing)\b", r"\butili[sz](?:e|es|ed|ing)\b",
    r"\bunderscor(?:e|es|ed|ing)\b", r"\billuminat(?:e|es|ed|ing)\b", r"\belevat(?:e|es|ed|ing)\b",
    r"\bfoster(?:s|ed|ing)?\b", r"\bembark(?:s|ed|ing)?\b", r"\brevolutioni[sz](?:e|es|ed|ing)\b",
    r"\bunlock(?:s|ed|ing)?\b", r"\bunveil(?:s|ed|ing)?\b", r"\bpropel(?:s|led|ling)?\b",
    r"\bbolster(?:s|ed|ing)?\b", r"\bspearhead(?:s|ed|ing)?\b", r"\bresonat(?:e|es|ed|ing)\b",
    r"\btransformative\b", r"\bmeticulous(?:ly)?\b", r"\bmultifaceted\b", r"\bvibrant\b",
    r"\bcutting-edge\b", r"\bparamount\b", r"\bseasoned\b", r"\btapestry\b", r"\btestament\b",
    r"\bsynerg(?:y|ies)\b", r"\bunderpinnings\b", r"\bparadigms?\b", r"\bexcited\b", r"\bthrilled\b",
    r"\bpassionate\b", r"\bdream role\b", r"\bimpactful\b", r"\binterplay\b", r"\bhoned\b",
    r"\bI share your (?:commitment|passion|values|vision)\b",
    r"\bharness(?:es|ed)?\s+(?:the power|our|their|my)\b", r"\bharnessing\b",
    # phrases
    r"\bproven track record\b", r"\bresults-driven\b", r"\bdetail-oriented\b", r"\bhard-working\b",
    r"\bstrategic thinker\b", r"\bdynamic leader\b", r"\bhoned my\b", r"\bthroughout my career\b",
    r"\bmy diverse background\b", r"\bit(?: is|'s) important to note\b", r"\bit(?: is|'s) worth mentioning\b",
    r"\bI was able to\b", r"\bI am confident\b", r"\bI'm confident\b", r"\bI believe (?:I would|my)\b",
    r"\b(?:great|perfect) fit\b", r"\baligns? perfectly\b", r"\bskills and expertise\b", r"\bbring my skills\b",
    r"\beager to\b", r"\bplease allow me\b", r"\bwriting to express\b", r"\bexpress my interest\b",
    r"\breaching out to express\b", r"\bcontribute to your team\b", r"\bI look forward to\b",
    r"\bthank you for considering\b", r"\bat your earliest convenience\b", r"\bwelcome the opportunity\b",
    r"\bcontinued success\b", r"\bin summary\b", r"\bin conclusion\b", r"\bin essence\b",
    r"\bsteep learning curve\b", r"\bin today's\b", r"\bimagine\b", r"\bpicture this\b",
]
# Tier 2: WARN everywhere. Honest uses exist; a warning left standing names the use.
SOFT_LEXICON = [
    r"\bnavigat(?:e|es|ed|ing)\b", r"\blandscapes?\b", r"\becosystems?\b", r"\brobust\b", r"\bseamless(?:ly)?\b",
    r"\bcomprehensive\b", r"\bcrucial\b", r"\benhanc(?:e|es|ed|ing)\b", r"\bfacilitat(?:e|es|ed|ing)\b",
    r"\binnovative\b", r"\bdynamic\b", r"\bnuanced\b", r"\bstreamlin(?:e|es|ed|ing)\b",
    r"\bdiverse\b", r"\bcross-functional collaboration\b", r"\bmeasurable outcomes\b", r"\bcontribute to\b",
]
CONTRACTIONS = re.compile(
    r"\b(?:I'm|I've|I'd|I'll|it's|that's|there's|what's|here's|let's|don't|doesn't|didn't|can't|couldn't|won't|"
    r"wouldn't|shouldn't|isn't|aren't|wasn't|weren't|hasn't|haven't|hadn't|we're|we've|we'd|we'll|they're|"
    r"they've|you're|you'll|you've)\b", re.I)
HOOK_OPENER_BAN = re.compile(r"^\s*(?:imagine\b|picture this\b|in today's\b|ever since I was\b|I have always been\b)", re.I)
HOOK_OPENER_WARN = re.compile(r"^\s*As an?\b", re.I)
WORK_AUTH_P1 = re.compile(r"\bvisa\b|\bresidence permit\b|\bwork permit\b|\bAufenthalt\w*|\bBlue Card\b", re.I)
WORK_AUTH_ANY = re.compile(r"\bvisa\b|\bresidence permit\b|\bwork permit\b|\bAufenthalt\w*|\bBlue Card\b|"
                           r"\bright to work\b|\bsponsorship\b", re.I)
BY_ABLE = re.compile(r"^\s*By\s+\w+ing\b[^.]{0,100},\s*(?:I|we)\b", re.I)

# Standing Approval Policy: no future-tense LEARNING promises (blocking).
# Approach statements ("In this role I would write...") are allowed and expected.
FUTURE_PROMISES = [
    r"\bi will (?:learn|master|get up to speed|close|build up|develop my)\b",
    r"\bi plan to (?:learn|complete|study|take)\b",
    r"\bget up to speed\b", r"\bquickly (?:learn|pick up|ramp)\b",
    r"\bi am (?:eager|keen) to learn\b", r"\bi would (?:learn|quickly)\b",
]

# CANDIDATE_STRATEGY Section 5B (blocking)
BANNED_CLAIMS = [
    (r"zero hallucination", "5B: 'zero hallucination'"),
    (r"\bAI Act (?:consulting|client|advisory work)", "5B: AI Act client work"),
    (r"30\+ EU (?:technology )?companies", "5B: EU-scoped client count"),
    (r"Curo[^.]{0,60}\b(?:launched|is live|went live|shipped)\b", "5B: Curo as launched/shipped"),
    (r"\b(?:launched|shipped)\s+Curo\b", "5B: Curo as launched/shipped"),
    (r"\b(?:delivered|shipped|launched|compressed)\b[^.]{0,30}16-week MVP", "5B: 16-week MVP is a scope, not a delivery"),
    (r"four years of legal", "5B: legal tenure beyond CV"),
    (r"\b20 ?(?:%|percent)[^.]{0,40}case[- ]prep", "5B: retired Khaitan figure"),
    (r"\bachieved\b[^.]{0,60}(?:Weekly Reciprocity|under 20 seconds|duplicate[- ]collection)", "5B: portfolio target as result"),
    (r"Article 35[^.]{0,60}\bclients?\b|\bclients?\b[^.]{0,60}Article 35", "5B: Article 35 client work"),
    (r"(?<!co-)\b(?:authored|wrote|led)\b[^.]{0,30}GIZ (?:policy )?brief", "5B: GIZ brief must read co-authored"),
    (r"(?:RasaHub|Typecast)[^.]{0,80}\b(?:user interviews|interviewed)\b", "5B: user research that did not happen"),
    (r"4\.3M\+?[^.]{0,60}\b(?:users|served|reached)\b", "5B: market figure as product metric"),
    (r"translat\w+ GDPR obligations[^.]{0,40}product changes", "5B: GDPR 'product changes' client claim"),
    # CMGRS team credit, PORTFOLIO_DOSSIER.md honesty rule 10
    (r"\bI\s+led\b[^.]{0,40}\b(?:redesign|rebuild|rebuilt|project|platform|turnaround)\b", "5B: CMGRS individual ownership ('I led...')"),
    (r"\b[Ll]ed\s+the\s+(?:redesign|rebuild|turnaround)\b", "5B: CMGRS individual ownership ('led the redesign')"),
    (r"\b[Tt]ook ownership of a failing\b", "5B: CMGRS individual ownership ('took ownership')"),
    (r"\bI\s+(?:rebuilt|redesigned|centralised|centralized|restructured|routed)\b[^.]{0,80}\b(?:grievance|case|CMGRS|IPAC|closure|sprint|requirements)\b", "5B: CMGRS team decision written as 'I'"),
    # north star never presented as the cause of the 28%
    (r"north star[^.]{0,120}\b(?:cut|cutting|reduc\w*|fell|drove|deliver\w*|led to)\b[^.]{0,60}\b2[08]\s?%", "5B: north star presented as the cause of the 28%"),
    (r"\b2[08]\s?%[^.]{0,80}\b(?:by|through|thanks to|because of)\b[^.]{0,40}north star", "5B: north star presented as the cause of the 28%"),
    # Curo repo audit, v1.4 (claude/CURO_BULLET_BANK.md sections 7 and 9). Curo was built solo with AI coding agents.
    (r"\b(?:CTO|co-?founder)\s+(?:leads|led|carries|carried)\b[^.]{0,30}\bbuild\b", "5B: Curo was built solo, no CTO-led build"),
    (r"\bco-built\b[^.]{0,40}\bCTO\b", "5B: Curo was built solo, no 'co-built with the CTO'"),
    (r"(?:cohort|patients?)[^.]{0,80}\bweekly feedback\b|\bweekly feedback\b[^.]{0,40}(?:cohort|patients?)", "5B: cohort 'weekly feedback' is unsourced (repo: user testing not started)"),
    (r"per-token cost model", "5B: no per-token cost model in the repo; write 'per-user daily AI budget'"),
    (r"\b38\s*-\s*47[- ]week\b|\b16-week MVP\b", "5B: the 38-47 week estimate and 16-week MVP scope are unsourced; use the 12-day tagged MVP until Vibhor confirms them"),
    (r"Charit[eé][^.]{0,40}\bpartner|\bpartner\w*\b[^.]{0,40}Charit[eé]", "5B: Charite BIH partnership unconfirmed; only with a real conversation or letter"),
    (r"\bshipped\s+(?:the\s+)?(?:Curo\s+)?MVP\b", "5B: Curo's MVP is a tagged build, not shipped; write 'built'"),
    (r"\bengine\b[^.]{0,40}\b(?:is live|went live|in production)\b", "5B: Curo's engine is deployed but not live or in production"),
    (r"\b(?:ML|machine[- ]learning)[- ]based (?:output )?(?:classifier|filter)\b", "5B: Curo's output filter is rule-based, not ML"),
    (r"MoSCoW[^.]{0,60}\bCuro\b|\bCuro\b[^.]{0,80}MoSCoW", "5B: no MoSCoW record in the Curo repo"),
]

BANNED_CLAIMS = [(rx.replace("Curo", PRODUCT), label) for rx, label in BANNED_CLAIMS]

# Retired figures, CANDIDATE_STRATEGY Section 5A/5B v9 and v9.2 (blocking)
RETIRED_FIGURES = [
    (r"\b30\s?%[^.\n]{0,60}\bresolution|\bresolution[^.\n]{0,60}\b30\s?%", "retired figure: 30% case resolution, use 28%"),
    (r"\b200\+?\s*(?:dormitory\s+|dorm\s+)?residents\b|200-person|200-resident", "retired figure: 200+ residents, use 109"),
    (r"\b30\+\s*advisory reports\b", "retired figure: 30+ advisory reports, use 14"),
    (r"\b15\s*-\s*20[- ](?:patient|person)\b|\b15 to 20 patients\b", "retired figure: 15-20 patient cohort"),
    (r"AI Governance in Practice|AI Governance Framework Evaluation", "retired thesis title, use 'From Guarantee to Guardrail...'"),
    (r"Operationalises Constitutional Equality", "thesis title wording: 'Operationalising Constitutional Equality'"),
    (r"\b(?:six|seven|6|7)\+?\s+years\b|~\s*7\s+years", "years phrasing: use '5+ years'"),
    (r"\b60\s*(?:-|to)\s*70\s*%", "retired figure: 60-70% false positive reduction, use '68% to 3% in simulation'"),
]

# Curo history that must not appear in any artefact (v1.5): FAIL in cv and cl, WARN in net.
LEGACY_CURO = [
    (r"\bOB-?GYN\b", "legacy Curo: OBGYN advisors"),
    (r"\bFlare Mode\b", "legacy Curo: Flare Mode"),
    (r"\bCuro Heroes\b", "legacy Curo: Curo Heroes"),
    (r"\bVan Westendorp\b", "legacy Curo: Van Westendorp study"),
    (r"\b15[- ]patients?\b|\b15-person\b|co-development cohort", "legacy Curo: the old 15-patient cohort"),
    (r"cohort[^.]{0,80}\bbefore any code\b|cohort[^.]{0,60}\bset\b[^.]{0,30}\bdecisions\b", "legacy Curo: a cohort that set decisions before code"),
    (r"\bArticle 50\b[^.]{0,60}\b(?:compliant|compliance|certified|sprint)\b|\bAI Act[- ]compliant\b", "Article 50 overclaim: write 'aligned to', never compliant or certified"),
    (r"\b59 of 61\b|\b61[- ]tasks?\b|\b61-task\b", "legacy Curo: 59 of 61 / 61 tasks"),
    (r"\bEU-region\b", "legacy Curo: EU-region providers"),
    (r"\bWeekly Reciprocity\b", "legacy Curo: Weekly Reciprocity Rate"),
    (r"\bDay[- ]21\b", "legacy Curo: the Day 21 promise"),
    (r"\b6-tap\b|\bsix-tap\b", "legacy Curo: the 6-tap budget"),
    (r"\bfive-event\b", "legacy Curo: the five-event chain"),
    (r"Curo[^.]{0,80}\b(?:PCOS|endometriosis|menstrual|women'?s health|German market|Germany)\b"
     r"|\b(?:PCOS|endometriosis)\b[^.]{0,80}\bCuro\b", "legacy Curo: an earlier condition or market"),
    (r"\bSai Prusni\b|\bPrusni\b", "Curo team: a named teammate"),
    (r"Berliner Startup|Stipendium", "Curo: the Stipendium claim is removed"),
    (r"\bformer co-?founder\b", "Curo team: a former co-founder"),
]
LEGACY_CURO = [(rx.replace("Curo", PRODUCT), label) for rx, label in LEGACY_CURO]
# Held claims and phrasing that needs a qualifier (v1.5): WARN only.
HELD_CURO = [
    (r"\bMDCG\b|\bDiGA\b", "held: MDCG and DiGA need a JD that screens them and Vibhor's confirmation"),
    (r"\bzero (?:unprotected|unlabelled)\b|\bzero health-data tables\b", "held: zero-gap audit claims wait for the Mumbai re-audit"),
    (r"\b(?:about )?43 test suites\b|\b49 (?:database )?migrations\b|\b50k[- ]user\b", "held: recount before using build counts"),
]
GATE_FIGURES = re.compile(r"\bat most 5\s?%|\bat least 60\s?%|\b100 red-team prompts\b|\bzero escapes\b", re.I)
COHORT = re.compile(r"\bcohort\b", re.I)
PILOT_RUN = re.compile(r"\b(?:ran|completed|conducted|finished)\s+(?:a|the|our)\s+(?:\w+\s+)?pilot\b", re.I)

# Curo figures that only exist in simulation (repo validation harness). Each needs "simulat" nearby.
SIMULATION_FIGURES = re.compile(r"\b68\s?%|\b19\.5\s?%|\b33\s*(?:-|to)\s*67\s?%")

CL_BUDGETS = [(95, 120), (115, 140), (105, 125), (40, 60)]          # APPLICATION_RULES.md Section 2, prose
CL_BUDGETS_TABLE = [(85, 110), (95, 125), (80, 115), (35, 55)]      # Section 2.5, evidence table
CL_TOTAL = (360, 400)
CL_TOTAL_TABLE = (300, 360)
TABLE_ROW = re.compile(r"^\s*\|(.+?)\|?\s*$")
SIGNOFF = re.compile(r"^\s*(?:Warm regards|With respect|Kind regards|Best regards|Sincerely|Regards)\b", re.I)
FIRST_PERSON = re.compile(r"\bI\b|\b(?:[Mm]e|[Mm]y|[Ww]e|[Oo]ur|[Uu]s)\b")   # v1.6: "We" at a sentence start counts
JARGON = ["BH-FDR", "Benjamini", "RAG", "MoSCoW", "PostHog", "classifier", "LightGBM", "SHAP", "Schrems", "Spearman",
          "permutation", "pgvector", "RLS", "case-crossover"]
KEY_FIGURES = [("28%", r"28%"), ("3M+", r"3M\+"), ("15,004", r"15,004"), ("16-week", r"16-week"),
               ("38-47", r"38-47"), ("68%", r"\b68%"), ("19.5%", r"19\.5%"), ("33-67%", r"33-67%"),
               ("12 days", r"\b12 days\b"), ("109", r"\b109\b"), ("25%", r"\b25%"),
               ("1.5M", r"1\.5M"), ("200+ cases", r"200\+ cases"), ("12M+", r"12M\+"),
               ("13 experiments", r"13 experiments"), ("about 35 patients", r"about 35 (?:patients|users)")]
STOP = set("""the and for with that this from into their your have has had are was were been being which what
when where while about over under than then they them there these those would could should will just only also
across around after before every other some more most such each very role team work product products company""".split())


LINKEDIN_URL = "linkedin.com/in/vibhorahlawat"

# Tools Vibhor has confirmed he works in (stated October 2026, plus tools already named in the master CVs and the
# Curo repo bank). Add a tool here only after he confirms it, in the same edit as CANDIDATE_STRATEGY.md Section 9B.
CONFIRMED_TOOLS = {
    "figma", "jira", "miro", "notion", "supabase", "posthog", "mixpanel", "amplitude", "sql", "postgresql",
    "postgres", "python", "r", "google cloud run", "cloud run", "google cloud", "react native", "google adk",
    "gemini", "cursor", "claude code", "expo", "aws bedrock", "amazon bedrock", "vertex ai", "google vertex ai",
    "pgvector",
}
JD_TOOL_CHECK = ["Figma", "Jira", "Miro", "Notion", "Supabase", "PostHog", "Mixpanel", "Amplitude", "SQL",
                 "Cursor", "Claude Code"]
LEVEL_WORDS = re.compile(r"\b(?:proficient|proficiency|expert|expertise|advanced|familiar|familiarity|basic|"
                         r"intermediate|beginner|working knowledge)\b", re.I)
TERM_STOP = {"and", "the", "for", "with", "into", "from"}
GENERIC_WORDS = {"product", "products", "development", "delivery", "management", "strategy", "planning",
                 "design", "process", "processes", "experience", "framework", "frameworks", "platform"}


def normalise(text):
    """Straighten curly apostrophes and quotes so every pattern sees one form."""
    return (text.replace("’", "'").replace("‘", "'")
                .replace("“", '"').replace("”", '"'))


def check_docx_links(path, fails):
    """The contact line's LinkedIn and portfolio must be clickable, not plain text."""
    import docx
    d = docx.Document(path)
    targets = [r.target_ref.lower() for r in d.part.rels.values() if "hyperlink" in r.reltype]
    if not any(LINKEDIN_URL in t for t in targets):
        fails.append("LinkedIn is not a clickable link on the contact line (set header.linkedin in the builder data)")
    if not any(CANONICAL_DOMAIN in t and RETIRED_DOMAIN not in t for t in targets):
        fails.append("portfolio vibhor.work is not a clickable link on the contact line (set header.portfolio)")


def docx_paragraphs(path):
    import docx
    d = docx.Document(path)
    paras = [p.text for p in d.paragraphs]
    tables = []
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                tables.append(c.text)
    headers = []
    for s in d.sections:
        headers += [p.text for p in s.header.paragraphs]
    return paras, tables, headers


def docx_body_items(path):
    """Body in document order: paragraph texts, with each table row written as '| left | right |'."""
    import docx
    from docx.text.paragraph import Paragraph
    from docx.table import Table
    d = docx.Document(path)
    items = []
    for el in d.element.body.iterchildren():
        tag = el.tag.split("}")[-1]
        if tag == "p":
            items.append(Paragraph(el, d).text)
        elif tag == "tbl":
            for row in Table(el, d).rows:
                cells = []
                for c in row.cells:
                    t = c.text.strip()
                    if t and (not cells or cells[-1] != t):   # merged cells repeat their text
                        cells.append(t)
                items.append("| " + " | ".join(cells) + " |")
    return items


def load(path, mode):
    """Return (full_text, body_text). body_text is the letter body for cl mode."""
    if path.endswith(".docx"):
        paras, tables, headers = docx_paragraphs(path)
        full = normalise("\n".join(headers + paras + tables))
        if mode == "cl":
            return full, normalise(letter_body(docx_body_items(path)))
        return full, full
    text = normalise(open(path, encoding="utf-8").read())
    if mode == "cl":
        return text, letter_body(text.split("\n"), joined=False, raw=text)
    return text, text


def letter_body(lines, joined=True, raw=None):
    """Cut the body from the line after 'Dear ...,' to the line before the sign-off.
    Falls back to the whole text when there is no salutation (bare draft paragraphs).
    In a .docx, items are joined with blank lines, except that table rows stay attached to the
    paragraph before and after them, so an evidence table and its lead-in form one paragraph."""
    start = next((i for i, l in enumerate(lines) if re.match(r"\s*Dear\b", l)), None)
    if start is None:
        return raw if raw is not None else "\n\n".join(l for l in lines if l.strip())
    end = next((i for i in range(start + 1, len(lines)) if SIGNOFF.match(lines[i])), len(lines))
    body = lines[start + 1:end]
    if not joined:
        return "\n".join(body)
    out, prev_row = "", False
    for l in body:
        l = l.strip()
        if not l:
            continue
        is_row = l.startswith("|")
        if out:
            out += "\n" if (is_row or prev_row) else "\n\n"
        out += l
        prev_row = is_row
    return out


def words(s):
    return len(re.findall(r"[A-Za-z0-9][A-Za-z0-9'%+.,/-]*", s))


def sentences(s):
    return [x.strip() for x in re.split(r"(?<=[.!?])\s+", s) if x.strip()]


def content_words(s):
    return {w.lower() for w in re.findall(r"[A-Za-z][A-Za-z-]{4,}", s)} - STOP


def split_rows(block):
    """Return (prose_without_rows, rows); rows are lists of cell strings."""
    prose, rows = [], []
    for line in block.split("\n"):
        m = TABLE_ROW.match(line)
        if m and line.strip().startswith("|"):
            rows.append([c.strip() for c in m.group(1).split("|") if c.strip()])
            if prose and prose[-1].rstrip().endswith(":"):
                prose[-1] = prose[-1].rstrip() + "."   # the lead-in ends where the table starts
        else:
            prose.append(line)
    return "\n".join(prose), rows


def check_links(text, fails):
    if RETIRED_DOMAIN in text:
        fails.append(f"retired domain '{RETIRED_DOMAIN}' used, write {CANONICAL_DOMAIN}")
    for link in DEEP_LINK.findall(text):
        if link.lower() not in ENTRY_LINKS:
            fails.append(f"deep link '{link}' not in the dossier entry list")
    if re.search(r"vibhor\.work/\S*umang", text, re.I):
        fails.append("Umang is never linked")
    if not GYRO_PAGE_LIVE and re.search(r"vibhor\.work/project/curo", text, re.I):
        fails.append("the product page link (vibhor.work/project/curo) is suspended until GYRO_PAGE_LIVE is True (use CMGRS, BeamerHeld or Haqdaar)")


def check_lexicon(text, fails, warns, mode):
    """APPLICATION_RULES.md Section 0.2a. The hard tier blocks letters and outreach; CVs only warn."""
    hard_sink = warns if mode == "cv" else fails
    seen = set()
    for pat in HARD_LEXICON:
        for m in re.finditer(pat, text, re.I):
            key = (m.group(0).lower(), "hard")
            if key in seen:
                continue
            seen.add(key)
            hard_sink.append(f"AI-isms lexicon, hard tier (0.2a): '{m.group(0)}', use the plain word")
    for pat in SOFT_LEXICON:
        for m in re.finditer(pat, text, re.I):
            key = (m.group(0).lower(), "soft")
            if key in seen:
                continue
            seen.add(key)
            warns.append(f"AI-isms lexicon, soft tier (0.2a): '{m.group(0)}', keep only if the plain word would be wrong")


def common(text, fails, warns, mode="cv"):
    check_links(text, fails)
    for ch, name in DASHES.items():
        for m in re.finditer(ch, text):
            ctx = text[max(0, m.start() - 30): m.end() + 30].replace("\n", " ")
            fails.append(f"{name}: ...{ctx}...")
    for pat in BANNED_PHRASES + FUTURE_PROMISES:
        for m in re.finditer(pat, text, re.I):
            fails.append(f"banned phrase '{m.group(0)}'")
    check_lexicon(text, fails, warns, mode)
    for pat, label in BANNED_CLAIMS + RETIRED_FIGURES:
        for m in re.finditer(pat, text, re.I):
            fails.append(f"{label}: '{m.group(0)[:70]}'")
    for m in SIMULATION_FIGURES.finditer(text):
        window = text[max(0, m.start() - 200): m.end() + 120].lower()
        if "simulat" not in window:
            fails.append(f"product figure '{m.group(0)}' without 'in simulation' (or 'simulated') nearby")
    for pat, label in LEGACY_CURO:
        for m in re.finditer(pat, text, re.I):
            (warns if mode == "net" else fails).append(f"{label}: '{m.group(0)[:60]}'")
    if not GYRO_COFOUNDER_ALLOWED and mode in ("cv", "cl"):
        for m in re.finditer(r"\bco-?found(?:er|ers|ed|ing)\b|\bCTO\b", text, re.I):
            fails.append(f"Gyro Health is run solo: no co-founder or CTO may be named ('{m.group(0)}')")
    for m in OLD_NAME.finditer(text):
        (warns if mode == "net" else fails).append("old product name 'Curo': the product is Gyro Health in every artefact (0.0)")
    for m in RENAME_MENTION.finditer(text):
        (warns if mode == "net" else fails).append(f"never mention that the product was renamed: '{m.group(0)[:50]}'")
    for pat, label in HELD_CURO:
        for m in re.finditer(pat, text, re.I):
            warns.append(f"{label}: '{m.group(0)[:50]}'")
    for m in GATE_FIGURES.finditer(text):
        window = text[max(0, m.start() - 160): m.end() + 120].lower()
        if not re.search(r"target|gate|plan|\bset\b|specif|designed|to be", window):
            warns.append(f"product gate figure without 'target', 'gate' or 'plan' nearby: '{m.group(0)}'")
    for m in COHORT.finditer(text):
        window = text[max(0, m.start() - 200): m.end() + 160].lower()
        if not re.search(r"pilot|planned|plan\b|about 35|from 11 nov|enrol", window):
            warns.append("a pilot cohort must read as the planned pilot cohort (add 'pilot', 'planned' or 'about 35')")
    for m in PILOT_RUN.finditer(text):
        (warns if mode == "net" else fails).append(f"Gyro Health has run no pilot: '{m.group(0)}'")
    for m in re.finditer(r"\bpilot (?:results|findings)\b|\binterview findings\b", text, re.I):
        warns.append(f"check: Gyro Health has no pilot results or interview findings yet: '{m.group(0)}'")
    for m in re.finditer(r"\b(?:steelman|expert|review) committees?\b|\bcommittee review\b", text, re.I):
        if "simulat" not in text[max(0, m.start() - 120): m.end() + 120].lower():
            warns.append(f"committee reviews are AI-composed: say 'simulated': '{m.group(0)}'")


def max_run_without_first_person(block):
    run = best = 0
    for s in sentences(block):
        if FIRST_PERSON.search(s):
            run = 0
        else:
            run += 1
            best = max(best, run)
    return best


def check_table(blocks, fails, warns):
    """APPLICATION_RULES.md Section 2.5. Returns True when the letter carries an evidence table."""
    located = [(i, rows) for i, b in enumerate(blocks) for rows in [split_rows(b)[1]] if rows]
    if not located:
        return False
    for i, rows in located:
        if i != 2:
            fails.append(f"evidence table found in Para {i + 1}; it belongs in Para 3 only (2.5)")
    rows = [r for _, rs in located for r in rs]
    if not 3 <= len(rows) <= 4:
        fails.append(f"evidence table has {len(rows)} rows, needs 3-4 (2.5)")
    seen_figs = {}
    for n, r in enumerate(rows, 1):
        if len(r) != 2:
            fails.append(f"evidence table row {n} needs exactly two cells, found {len(r)}: {' | '.join(r)[:60]}")
            continue
        left, right = r
        if words(left) > 6:
            warns.append(f"evidence table row {n}: left cell {words(left)} words, cap 6 (the JD's own phrase)")
        if words(right) > 25:
            warns.append(f"evidence table row {n}: right cell {words(right)} words, cap 25")
        if not re.search(r"\d", right):
            fails.append(f"evidence table row {n}: right cell has no numeral (one specific time, with a number): {right[:60]}")
        for fig in re.findall(r"\d[\d,.]*(?:\s?%|\+|M\+?|k\b)?", right):
            f = fig.strip().rstrip(".,")
            if re.search(r"%|\+|M", f) or len(re.sub(r"\D", "", f)) >= 2:
                if f in seen_figs and seen_figs[f] != n:
                    fails.append(f"evidence table: figure '{f}' appears in rows {seen_figs[f]} and {n} (2.5: no figure twice)")
                seen_figs.setdefault(f, n)
    return True


def check_cl(text, body, fails, warns, jd):
    blocks = [b.strip() for b in re.split(r"\n\s*\n", body) if words(b) >= 15]
    is_table = check_table(blocks, fails, warns)
    lo_t, hi_t = CL_TOTAL_TABLE if is_table else CL_TOTAL
    budgets = CL_BUDGETS_TABLE if is_table else CL_BUDGETS
    total = sum(words(b) for b in blocks)
    if not lo_t <= total <= hi_t:
        fails.append(f"word count {total}, must be {lo_t}-{hi_t} ({'evidence table' if is_table else 'prose'} letter)")
    prose_blocks = [split_rows(b)[0] for b in blocks]
    if len(blocks) != 4:
        fails.append(f"found {len(blocks)} body paragraphs, expected 4")
    else:
        p1, p2, p3, p4 = blocks
        for i, (b, (lo, hi)) in enumerate(zip(blocks, budgets), 1):
            n = words(b)
            if not lo <= n <= hi:
                warns.append(f"Para {i}: {n} words, budget {lo}-{hi}")
        if GYRO_COFOUNDER_ALLOWED:
            has_disclosure = re.search(r"co-?found", p2 + p3, re.I)
        else:
            has_disclosure = FOUNDER_DISCLOSURE.search(p2 + p3)
        if not has_disclosure:
            fails.append("founder disclosure not found in Para 2 or 3 (name the founder role, e.g. 'the health app I founded')")
        disclosures = len(FOUNDER_DISCLOSURE.findall(body))
        if disclosures > 1:
            warns.append(f"founder disclosure appears {disclosures} times, it should be stated once")
        links = DEEP_LINK.findall(body)
        if len(links) != 1:
            fails.append(f"{len(links)} portfolio deep links in the body, need exactly 1")
        elif not DEEP_LINK.search(p3):
            fails.append("portfolio deep link is not in Para 3")
        elif is_table and any(DEEP_LINK.search(" ".join(r)) for r in split_rows(p3)[1]):
            fails.append("portfolio deep link sits inside the evidence table; it belongs in the prose (2.5)")
        if re.match(r"\s*(?:I am|I'm|I have)\b", p1):
            fails.append("Para 1 opens on the applicant, not the reader's world (2.1)")
        # v1.6 hook checks
        p1_sents = sentences(prose_blocks[0])
        first = p1_sents[0] if p1_sents else ""
        if HOOK_OPENER_BAN.match(first):
            fails.append(f"Para 1 opens on a banned hook ('Imagine', 'Picture this', 'In today's'...): {first[:60]}")
        if first.rstrip().endswith("?"):
            fails.append("Para 1 opens on a rhetorical question; open on the reader's world (2.1)")
        if HOOK_OPENER_WARN.match(first):
            warns.append(f"Para 1 opens on 'As a...', which usually leads with his CV; the hook comes from the reader's world (2.1): {first[:60]}")
        if words(first) > 25:
            warns.append(f"first sentence {words(first)} words; the hook lands harder under 25")
        for m in WORK_AUTH_P1.finditer(p1):
            fails.append(f"work-authorisation detail in Para 1 ('{m.group(0)}'): never in a hook (Story Bank, P4)")
        for b in (p2, p3, p4):
            for m in WORK_AUTH_ANY.finditer(b):
                warns.append(f"work-authorisation detail in the letter ('{m.group(0)}'): recruiter outreach only unless the JD makes it relevant (Objection 2)")
        # v1.2 spine and forward-move heuristics
        if not re.search(r"\bI would\b|\bI'd\b", p3):
            warns.append("Para 3 has no 'I would' approach sentence (forward move, Section 2)")
        overlap = content_words(p1) & content_words(p4)
        if len(overlap) < 2:
            warns.append(f"Para 4 shares {len(overlap)} content words with Para 1; the close should return to the thesis")
        for i, b in ((1, prose_blocks[0]), (3, prose_blocks[2])):
            run = max_run_without_first_person(b)
            if run >= 3:
                warns.append(f"Para {i}: {run} consecutive sentences with no I or we; intel may not be tied to his evidence")
        if re.search(r"@|\+49|\bcan be reached\b", p4):
            warns.append("Para 4 carries contact details; they belong in the header")
    for b in prose_blocks:
        for q in re.findall(r"[^.?!]*\?", b):
            fails.append(f"question in the letter body (Para 4 and every other paragraph close on statements): {q.strip()[:60]}")
    if "\n- " in body or "•" in body:
        fails.append("bullet points found in cover letter (an evidence table is the only list form allowed, 2.5)")
    sents = [s for b in prose_blocks for s in sentences(b)]
    for s in sents:
        if words(s) > 30:
            warns.append(f"sentence over 30 words ({words(s)}): {s[:60]}...")
        if BY_ABLE.match(s):
            warns.append(f"'By [doing X], I...' formula (0.2): {s[:60]}...")
    if sents and not any(words(s) < 10 for s in sents):
        warns.append("no sentence under 10 words, rhythm check fails")
    lengths = [words(s) for s in sents]
    if len(lengths) >= 4 and statistics.pstdev(lengths) < 5:
        warns.append(f"sentence lengths barely vary (standard deviation {statistics.pstdev(lengths):.1f} words, want 5+): mix long and short")
    n_contr = len(CONTRACTIONS.findall(body))
    if n_contr == 0:
        warns.append("no contractions; Tracks A, D and E want 1-2 (I'd, it's, that's), 3 at most (0.3a); may stand on B and C as 'formal register'")
    elif n_contr > 3:
        warns.append(f"{n_contr} contractions; minimise them: 1-2, 3 at most (0.3a)")


def find_summary(lines):
    for i, l in enumerate(lines):
        if l.strip().upper() == "PROFILE SUMMARY":
            for j in range(i + 1, min(i + 4, len(lines))):
                if lines[j].strip():
                    return j, lines[j].strip()
    return None, None


def count_bullets_after(lines, anchor_re):
    """Count bullet-like lines under the first org line matching anchor_re."""
    idx = next((i for i, l in enumerate(lines) if re.match(anchor_re, l.strip())), None)
    if idx is None:
        return None
    n, i, skipped = 0, idx + 1, 0
    while i < len(lines):
        l = lines[i].strip()
        if not l:
            i += 1
            continue
        is_bullet = bool(re.match(r"(?:[•*-]|\d+\.)\s", l)) or words(l) >= 12
        if is_bullet and not l.startswith("*("):
            n += 1
        elif skipped < 2 and n == 0:
            skipped += 1          # role line under the org line
        else:
            break
        i += 1
    return n


def check_cv(text, fails, warns, jd):
    lines = text.splitlines()
    if not re.search(r"Portfolio:\s*vibhor\.work", text):
        fails.append("contact line missing 'Portfolio: vibhor.work'")
    s_idx, summary = find_summary(lines)
    if summary is None:
        fails.append("no PROFILE SUMMARY section found")
    else:
        n = words(summary)
        if n > 55:
            fails.append(f"Profile Summary {n} words, cap 55")
        if re.search(r"\bI\b", summary):
            fails.append("Profile Summary uses 'I'")
        if len(sentences(summary)) < 3:
            warns.append(f"Profile Summary has {len(sentences(summary))} sentence(s); spec is three full sentences")
    jd_text = (jd or "").lower()
    for i, line in enumerate(lines):
        if i == s_idx:
            continue
        l = line.strip()
        n = words(l)
        if re.match(r"(?:\*\*)?(?:Problem|Solution):", l):
            if n > 40:
                warns.append(f"{l.split(':')[0]} line over 40 words ({n}): {l[:60]}...")
            continue
        if re.match(r"(?:[•*-]|\d+\.)\s", l) or 20 <= n <= 60:
            if n > 35:
                warns.append(f"bullet over 35 words ({n}): {l[:60]}...")
            for term in JARGON:
                pat = r"\b" + re.escape(term) + r"\b"
                if re.search(pat, l) and term.lower() not in jd_text:
                    warns.append(f"method name '{term}' not in the JD{'' if jd else ' (no --jd given)'}: {l[:50]}...")
    curo = count_bullets_after(lines, r"(?:#+\s*)?(?:Gyro Health|Curo)\b")
    if curo is not None and curo > 5:
        warns.append(f"the product entry has {curo} bullets, cap 5")
    ipac = count_bullets_after(lines, r"(?:#+\s*)?Indian Political Action Committee")
    if ipac is not None and ipac > 4:
        warns.append(f"IPAC has {ipac} bullets, cap 4")
    body = "\n".join(lines)
    for label, fig in KEY_FIGURES:
        c = len(re.findall(fig, body))
        if c > 2:
            warns.append(f"figure '{label}' used {c} times, cap 2 (Section 11.1)")
    if re.search(r"\b(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec) \d{4} to (?:Present|Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)", text):
        warns.append("old date format 'Mon YYYY to ...'; use 'Mon YYYY - Mon YYYY'")
    if re.search(r"^\s*(?:##\s*)?PRODUCTS BUILT\s*$", text, re.M | re.I):
        warns.append("'Products Built' heading; renamed 'Side Projects' in v12")


def _top_level_split(s, seps=","):
    """Split on separators that sit outside brackets."""
    out, depth, cur = [], 0, ""
    for ch in s:
        if ch in "([":
            depth += 1
        elif ch in ")]":
            depth = max(0, depth - 1)
        if ch in seps and depth == 0:
            out.append(cur)
            cur = ""
        else:
            cur += ch
    out.append(cur)
    return [x.strip() for x in out if x.strip()]


def cv_structure(path, text):
    """Return (aoe_terms, tools_line_or_None, echo_body). Reads .docx or .md/.txt."""
    if path.endswith(".docx"):
        import docx
        d = docx.Document(path)
        terms, seen = [], set()
        for t in d.tables:
            for row in t.rows:
                for c in row.cells:
                    v = re.sub(r"^[\s•*\-]+", "", c.text).strip()
                    if v and v not in seen:
                        seen.add(v)
                        terms.append(v)
        tools, body = None, []
        for p in d.paragraphs:
            t = p.text.strip()
            if re.match(r"Tools:", t):
                tools = t
            else:
                body.append(t)
        return terms, tools, "\n".join(body)
    lines = text.splitlines()
    start = next((i for i, l in enumerate(lines) if re.match(r"^\W*(?:#+\s*)?AREAS OF EXPERTISE\W*$", l.strip(), re.I)), None)
    if start is None:
        return None, None, text
    end = len(lines)
    for j in range(start + 1, len(lines)):
        if re.match(r"^\W*(?:#+\s*)?(?:KEY ACHIEVEMENTS|PROFESSIONAL EXPERIENCE|SIDE PROJECTS|CERTIFICATIONS|EDUCATION)\W*$",
                    lines[j].strip(), re.I):
            end = j
            break
    block = lines[start + 1:end]
    terms, tools = [], None
    for l in block:
        t = re.sub(r"^[\s•*\-]+|[\s*]+$", "", l).replace("**", "").strip()
        if not t or t.startswith("(") and t.endswith(")"):
            continue
        if re.match(r"Tools:", t):
            tools = t
            continue
        for chunk in re.split(r"\t|\||·|•|;", t):
            terms.extend(_top_level_split(chunk))
    body = "\n".join(lines[:start] + lines[end:])
    return terms, tools, body


def _stem_pattern(word):
    """Regex fragment that matches a word and its common forms in the body."""
    w = word.lower()
    special = {"leadership": r"\b(?:lead|led|leader)", "management": r"\bmanag", "analytics": r"\banalyt",
               "analysis": r"\banaly", "prioritisation": r"\bpriorit", "prioritization": r"\bpriorit",
               "roadmapping": r"\broadmap", "compliance": r"\bcompl(?:ian|y)", "governance": r"\bgovern",
               "delivery": r"\bdeliver", "evaluation": r"\bevaluat", "architecture": r"\barchitect",
               "research": r"\bresearch", "discovery": r"\bdiscover", "alignment": r"\balign",
               "stakeholder": r"\bstakeholder", "engagement": r"\bengag", "regulatory": r"\bregulat"}
    if w in special:
        return special[w]
    if len(w) >= 8:
        return r"\b" + re.escape(w[:6])
    if len(w) >= 5:
        return r"\b" + re.escape(w[:5])
    return r"\b" + re.escape(w) + r"\b"


def _echo_status(part, body):
    """'ok' or 'none' for one term part. Distinctive words must echo; generic words only count when a part has nothing else."""
    toks = re.findall(r"[A-Za-z][A-Za-z0-9+.]*(?:-[A-Za-z0-9+.]+)*", part)
    sig = []
    for tk in toks:
        if tk.lower() in TERM_STOP:
            continue
        if tk.isupper() and 2 <= len(tk) <= 5:
            sig.append((tk, r"\b" + re.escape(tk) + r"\b", False))
        elif "-" in tk:
            parts = [x for x in tk.split("-") if x]
            sig.append((tk, r"[- ]?".join(_stem_pattern(x) for x in parts), False))
        elif len(tk) >= 5:
            sig.append((tk, _stem_pattern(tk), tk.lower() in GENERIC_WORDS))
    if not sig:
        return "ok"
    real = [s for s in sig if not s[2]]
    pool = real if real else sig        # a part with a distinctive word must echo that word, not a generic one
    return "ok" if any(re.search(pat, body, re.I) for _, pat, _g in pool) else "none"


def check_aoe_and_tools(path, text, fails, warns, jd):
    terms, tools, body = cv_structure(path, text)
    if terms is None:
        fails.append("no AREAS OF EXPERTISE section found")
        return
    if len(terms) < 6:
        warns.append(f"Areas of Expertise lists {len(terms)} terms; the 3x3 grid expects 9 (minimum 6)")
    for term in terms:
        top = re.sub(r"\([^)]*\)", " ", term)
        parts = [x.strip() for x in re.split(r"\s+and\s+|\s*&\s*|\s*/\s*|,\s*", top) if x.strip()]
        for part in parts:
            st = _echo_status(part, body)
            if st == "none":
                fails.append(f"Areas of Expertise term '{term}': no echo of '{part}' in any bullet, Key Achievement or the summary "
                             f"(11.3a: reword the closest bullet into this word, or drop the term; a RED term never goes on the CV)")
        for inner in re.findall(r"\(([^)]*)\)", term):
            for item in [x.strip() for x in inner.split(",") if x.strip()]:
                if item.lower() in CONFIRMED_TOOLS:
                    continue
                if _echo_status(item, body) == "none":
                    warns.append(f"Areas of Expertise term '{term}': bracketed '{item}' has no echo in the body")
    if tools is None:
        fails.append("no Tools line (11.3b): add 'Tools: ...' with confirmed tools only, directly under Areas of Expertise")
        return
    listed = [x.strip() for x in _top_level_split(re.sub(r"^Tools:\s*", "", tools), ",;") if x.strip()]
    if LEVEL_WORDS.search(tools):
        fails.append("Tools line carries a skill-level word (proficient, expert, advanced, familiar...): names only")
    for t in listed:
        if t.lower() not in CONFIRMED_TOOLS:
            fails.append(f"Tools line: '{t}' is not a tool Vibhor has confirmed (CANDIDATE_STRATEGY.md Section 9B); remove it")
    if len(listed) > 10:
        warns.append(f"Tools line lists {len(listed)} tools, cap 10; keep the JD's tools first and trim the rest")
    if jd:
        have = {t.lower() for t in listed}
        for t in JD_TOOL_CHECK:
            if re.search(r"\b" + re.escape(t) + r"\b", jd, re.I) and t.lower() not in have:
                warns.append(f"the JD names {t}, a confirmed tool; put it on the Tools line")


HM_ITEMS = 12
HM_LINE = re.compile(r"^\s*(?:[-*]\s*)?(\d{1,2})[.)]\s+(.*)$")


def check_hm(path):
    """The Hiring Manager Review record. One line per rubric item:
    '<n>. <item name> | PASS or FAIL[, revised] | "quoted line from the draft"'."""
    fails, warns, seen, passed = [], [], {}, 0
    for line in normalise(open(path, encoding="utf-8").read()).split("\n"):
        m = HM_LINE.match(line)
        if not m:
            continue
        n, rest = int(m.group(1)), m.group(2)
        if not 1 <= n <= HM_ITEMS:
            continue
        if n in seen:
            fails.append(f"item {n} appears twice")
            continue
        seen[n] = rest
        verdict = re.search(r"\b(PASS|FAIL)\b", rest)
        if not verdict:
            fails.append(f"item {n}: no PASS or FAIL")
        elif verdict.group(1) == "PASS":
            passed += 1
        elif not re.search(r"revised", rest, re.I):
            warns.append(f"item {n}: FAIL not marked revised; name it in the decisions log")
        quote = re.search(r'"([^"]+)"', rest)
        if not quote or len(quote.group(1).split()) < 3:
            fails.append(f"item {n}: no quoted line from the draft as evidence (3+ words in double quotes)")
    missing = [str(i) for i in range(1, HM_ITEMS + 1) if i not in seen]
    if missing:
        fails.append(f"rubric items missing: {', '.join(missing)} (APPLICATION_RULES.md 8A has {HM_ITEMS})")
    print(f"QA HM {path}: {'FAIL' if fails else 'PASS'} ({passed}/{HM_ITEMS} pass)")
    for f in fails: print(f"  FAIL  {f}")
    for w in warns: print(f"  WARN  {w}")
    sys.exit(1 if fails else 0)


def main():
    args = sys.argv[1:]
    if len(args) == 2 and args[0] == "hm":
        check_hm(args[1])
    jd_path = None
    if "--jd" in args:
        k = args.index("--jd")
        jd_path = args[k + 1] if k + 1 < len(args) else None
        args = args[:k] + args[k + 2:]
    if len(args) != 2 or args[0] not in ("cl", "cv", "net"):
        print(__doc__); sys.exit(2)
    mode, path = args
    jd = open(jd_path, encoding="utf-8").read() if jd_path else None
    full, body = load(path, mode)
    fails, warns = [], []
    common(full, fails, warns, mode)
    if mode in ("cl", "cv") and path.endswith(".docx"):
        check_docx_links(path, fails)
    if mode == "cl":
        check_cl(full, body, fails, warns, jd)
    elif mode == "cv":
        check_cv(full, fails, warns, jd)
        check_aoe_and_tools(path, full if not path.endswith(".docx") else "", fails, warns, jd)
    print(f"QA {mode.upper()} {path}: {'FAIL' if fails else 'PASS'}")
    for f in fails: print(f"  FAIL  {f}")
    for w in warns: print(f"  WARN  {w}")
    sys.exit(1 if fails else 0)


if __name__ == "__main__":
    main()
